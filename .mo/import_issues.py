#!/usr/bin/env python3
"""One-off import of the Group 7 task board into GitHub issues (no ongoing sync).

Creates one issue per Task ID, titled ``[ID] Title``, with the board's labels,
gate milestone (G0-G5), assignee and a body in the shape of the team's task
issue template. Dependencies become issue links, so each dependency's issue
also shows "mentioned in #N".

Usage (from the repo root, on your machine)
-------------------------------------------
  python .mo/import_issues.py BOARD.xlsx --dry-run          # offline preview, writes .mo/issues_preview.md
  python .mo/import_issues.py BOARD.xlsx --limit 3          # create the first 3, check them on GitHub
  python .mo/import_issues.py BOARD.xlsx                    # create the rest (existing IDs are skipped)
  python .mo/import_issues.py BOARD.xlsx --done RV-01       # also close these as completed

Reading an .xlsx needs openpyxl (``python -m pip install openpyxl``). A CSV
export of the Tasks sheet works without it.

Safe to rerun: an issue whose title starts with ``[ID]`` (open or closed) is
skipped, and pull requests are ignored. Existing issues are never edited.
Issue numbers are appended to ``.mo/issue_numbers.csv`` after every issue.
``deliver.py`` reads it to add "Closes #N" to PRs, and the numbers can be
pasted into the board's "GitHub Issue #" column by Task ID.

Auth, first that works: GH_TOKEN / GITHUB_TOKEN, ``gh auth token``, or the
token Git already stores for github.com (``git credential fill``). Issues are
created as whoever owns that token.

Only the standard library is needed apart from openpyxl for .xlsx input.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import re
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

HERE = Path(__file__).resolve().parent
TEAM_SLUG = "USYD-ELEC5308/26-S2-Assignment2-Group7"
API_URL = "https://api.github.com"

# Confirmed handles (A1 merge commits and A2 pull requests; docs/repository_setup.md).
# The board's Team sheet overrides these when its GitHub column is filled in.
HANDLES = {
    "Mohammad": "kkaaman",
    "Alan": "at0941",
    "Abdullah": "Abdullah-Almutairi-uni",
    "Zhaolin": "Zhaolin-Huang",
    "Ricky": "RickyPurani",
    "Saviour": "skwo0237",
}

# Used only if the workbook has no Milestones / Labels sheet (e.g. CSV input).
GATES = {
    "G0": ("Contracts and runnable skeleton", "2026-10-14"),
    "G1": ("Navigation gate", "2026-10-18"),
    "G2": ("First obstacle stop", "2026-10-21"),
    "G3": ("Baseline freeze", "2026-10-25"),
    "G4": ("Extension admission", "2026-10-28"),
    "G5": ("Submission ready", "2026-10-31"),
}
LABEL_COLOURS = {"carla": "0F766E", "category": "0F766E", "extension": "7C3AED",
                 "priority:p0": "B91C1C", "priority:p1": "D97706", "priority:p2": "64748B",
                 "workstream": "1D4ED8"}

TITLE_ID = re.compile(r"^\[([A-Z0-9][A-Z0-9-]*)\]")
MISSING = object()  # GitHub answered 404 where that is an expected answer


# --------------------------------------------------------------------------- board
def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _date(value) -> dt.date | None:
    if value in (None, ""):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return dt.date.fromisoformat(str(value).strip()[:10])


def _split(value, sep: str) -> list[str]:
    return [part.strip() for part in _text(value).split(sep) if part.strip()]


def read_board(path: Path) -> dict:
    """Return tasks plus the Labels, Milestones and Team sheets when present."""
    board = {"tasks": [], "labels": {}, "milestones": {}, "handles": {}}
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as fh:
            board["tasks"] = [dict(row) for row in csv.DictReader(fh)]
        return board
    try:
        import openpyxl
    except ImportError:
        sys.exit("Reading .xlsx needs openpyxl:  python -m pip install openpyxl")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    def rows(sheet: str) -> list[dict]:
        if sheet not in wb.sheetnames:
            return []
        it = wb[sheet].iter_rows(values_only=True)
        header = [_text(h) for h in next(it)]
        return [dict(zip(header, r)) for r in it if r and r[0] not in (None, "")]

    board["tasks"] = rows("Tasks")
    for row in rows("Labels"):
        board["labels"][_text(row.get("Name"))] = (
            _text(row.get("Colour")) or "0F766E", _text(row.get("Description")))
    for row in rows("Milestones"):
        board["milestones"][_text(row.get("Gate"))] = (
            _text(row.get("Title")), _date(row.get("Due Date")), _text(row.get("Exit Criteria")))
    for row in rows("Team"):
        if _text(row.get("GitHub")):
            board["handles"][_text(row.get("Name"))] = _text(row.get("GitHub"))
    return board


def normalise(raw: dict) -> dict:
    """Map board column names (or the private tasks.csv names) to one shape."""
    def get(*names):
        for name in names:
            if name in raw and raw[name] not in (None, ""):
                return raw[name]
        return None

    return {
        "id": _text(get("Task ID", "id")),
        "title": _text(get("Title", "title")),
        "description": _text(get("Description", "description")),
        "acceptance": _text(get("Acceptance Criteria", "acceptance")),
        "category": _text(get("Category", "category")),
        "workstream": _text(get("Workstream", "workstream")),
        "gate": _text(get("Milestone", "milestone")),
        "priority": _text(get("Priority", "priority")),
        "owner": _text(get("Assignee", "owner")),
        "owner_handle": _text(get("Assignee GitHub")),
        "reviewer": _text(get("Reviewer", "reviewer")),
        "start": _date(get("Start Date", "start")),
        "due": _date(get("Due Date", "due")),
        "hours": _text(get("Estimate (h)", "hours")),
        "depends": [d.strip() for d in _text(get("Depends On", "depends_on")).split(",") if d.strip()],
        "carla": _text(get("CARLA Required", "carla")) or "N",
        "files": _split(get("Module / Files", "files"), ";"),
        "evidence": _text(get("Evidence Produced", "evidence")),
        "report": _text(get("Report Section")),
        "labels": _split(get("Labels"), ";"),
        "status": _text(get("Status")) or "Todo",
        "issue": _text(get("GitHub Issue #")),
        "notes": _text(get("Notes / Assumptions")),
    }


def derive_labels(task: dict) -> list[str]:
    """Labels for CSV input that has no Labels column (same scheme as the board)."""
    out = []
    ws = task["workstream"].lower()
    if ws.startswith("e") and ws[1:].isdigit():
        out.append(f"extension:{ws}")
    elif ws:
        out.append(f"workstream:{ws}")
    if task["priority"]:
        out.append(f"priority:{task['priority'].split()[0].lower()}")
    if task["carla"].upper() == "Y":
        out.append("carla-required")
    if task["category"]:
        out.append("category:" + re.sub(r"[^a-z0-9]+", "-", task["category"].lower()).strip("-"))
    return out


def topological(tasks: list[dict]) -> list[dict]:
    """Dependencies first (board order otherwise), so every dependency link resolves."""
    by_id = {t["id"]: t for t in tasks}
    order = {t["id"]: i for i, t in enumerate(tasks)}
    pending = {t["id"]: {d for d in t["depends"] if d in by_id} for t in tasks}
    done: list[dict] = []
    ready = sorted((i for i, deps in pending.items() if not deps), key=order.get)
    placed: set[str] = set()
    while ready:
        tid = ready.pop(0)
        placed.add(tid)
        done.append(by_id[tid])
        newly = [i for i, deps in pending.items()
                 if i not in placed and i not in ready and deps <= placed]
        ready = sorted(set(ready) | set(newly), key=order.get)
    leftovers = [t for t in tasks if t["id"] not in placed]  # cycles: keep board order
    return done + leftovers


# --------------------------------------------------------------------------- body
def _day(d: dt.date | None) -> str:
    return f"{d:%a} {d.day} {d:%b}" if d else "?"


def person(name: str, handles: dict) -> str:
    # No @: a mention in 223 bodies would notify people on every issue.
    handle = handles.get(name)
    return f"{name} ({handle})" if handle else (name or "unassigned")


def issue_body(task: dict, handles: dict, numbers: dict, imported_on: str) -> str:
    deps = []
    for dep in task["depends"]:
        deps.append(f"#{numbers[dep]} {dep}" if dep in numbers else dep)
    hours = f" (estimate {task['hours']} h)" if task["hours"] else ""
    lines = [
        f"**Task ID:** {task['id']}",
        f"**Owner / reviewer:** {person(task['owner'], handles)} / {person(task['reviewer'], handles)}",
        f"**Start → due:** {_day(task['start'])} → {_day(task['due'])}{hours}",
        f"**Depends on:** {', '.join(deps) if deps else 'none'}",
        f"**CARLA required:** {task['carla']}",
        f"**Gate:** {task['gate']} · **Priority:** {task['priority']} · "
        f"**Category:** {task['category']} · **Workstream:** {task['workstream']}",
        "",
        "### Description",
        "",
        task["description"] or "_None on the board._",
        "",
        "### Acceptance criteria",
        "",
    ]
    criteria = _split(task["acceptance"], ";") or ["_None on the board._"]
    lines += [f"- [ ] {c[:1].upper()}{c[1:].rstrip('.')}" for c in criteria]
    lines += ["", "### Module / files", ""]
    lines += [", ".join(f"`{f}`" for f in task["files"]) if task["files"] else "_None listed._"]
    lines += ["", "### Evidence produced", "", task["evidence"] or "_None listed._"]
    if task["report"]:
        lines += ["", f"**Report section:** {task['report']}"]
    if task["notes"]:
        lines += ["", f"> {task['notes']}"]
    lines += ["", f"<sub>Imported from the Group 7 task board on {imported_on}. "
              "Dates are not synced; the board stays the planning record.</sub>"]
    return "\n".join(lines) + "\n"


# --------------------------------------------------------------------------- GitHub
def find_token() -> str:
    for var in ("GH_TOKEN", "GITHUB_TOKEN"):
        if os.environ.get(var):
            return os.environ[var].strip()
    try:
        out = subprocess.run(["gh", "auth", "token"], capture_output=True, text=True, timeout=20)
        if out.returncode == 0 and out.stdout.strip():
            return out.stdout.strip()
    except (OSError, subprocess.SubprocessError):
        pass
    try:  # Git Credential Manager keeps an OAuth token for github.com after a push
        out = subprocess.run(
            ["git", "credential", "fill"], input="protocol=https\nhost=github.com\n\n",
            capture_output=True, text=True, timeout=60,
            env={**os.environ, "GIT_TERMINAL_PROMPT": "0"})
        for line in out.stdout.splitlines():
            if line.startswith("password="):
                return line.split("=", 1)[1].strip()
    except (OSError, subprocess.SubprocessError):
        pass
    sys.exit("No GitHub token found. Set GH_TOKEN, run `gh auth login`, or push once with git "
             "so Git Credential Manager stores a token.")


class GitHub:
    def __init__(self, repo: str, token: str, api_url: str, delay: float):
        self.base = f"{api_url.rstrip('/')}/repos/{repo}"
        self.token = token
        self.delay = delay
        self.writes = 0

    def call(self, method: str, path: str, payload: dict | None = None,
             ok_missing: bool = False):
        url = path if path.startswith("http") else self.base + path
        data = json.dumps(payload).encode("utf-8") if payload is not None else None
        for attempt in range(6):
            req = urllib.request.Request(url, data=data, method=method, headers={
                "Authorization": f"Bearer {self.token}",
                "Accept": "application/vnd.github+json",
                "X-GitHub-Api-Version": "2022-11-28",
                "User-Agent": "elec5308-g7-issue-import",
                **({"Content-Type": "application/json"} if data else {}),
            })
            try:
                with urllib.request.urlopen(req, timeout=60) as resp:
                    body = resp.read()
                    if method != "GET":
                        self.writes += 1
                        time.sleep(self.delay)  # GitHub asks for spacing between writes
                    return (json.loads(body) if body else None), resp.headers
            except urllib.error.HTTPError as err:
                text = err.read().decode("utf-8", "replace")
                if err.code == 404 and ok_missing:
                    return MISSING, err.headers
                limited = err.code == 429 or (err.code == 403 and (
                    "rate limit" in text.lower() or err.headers.get("retry-after")
                    or err.headers.get("x-ratelimit-remaining") == "0"))
                if limited and attempt < 5:
                    wait = int(err.headers.get("retry-after") or 0)
                    reset = err.headers.get("x-ratelimit-reset")
                    if not wait and reset and err.headers.get("x-ratelimit-remaining") == "0":
                        wait = max(1, int(reset) - int(time.time()) + 1)
                    wait = wait or 60 * (attempt + 1)
                    print(f"  rate limited by GitHub; waiting {wait} s")
                    time.sleep(wait)
                    continue
                raise SystemExit(f"GitHub {method} {url} failed: HTTP {err.code}\n{text}") from None
        raise SystemExit(f"GitHub {method} {url}: still rate limited after retries")

    def paged(self, path: str) -> list:
        items, url = [], self.base + path + ("&" if "?" in path else "?") + "per_page=100"
        while url:
            page, headers = self.call("GET", url)
            items += page or []
            match = re.search(r'<([^>]+)>;\s*rel="next"', headers.get("link") or "")
            url = match.group(1) if match else None
        return items


# --------------------------------------------------------------------------- main
def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("board", type=Path, help="task board .xlsx (or a CSV export of the Tasks sheet)")
    p.add_argument("--repo", default=TEAM_SLUG)
    p.add_argument("--dry-run", action="store_true", help="no GitHub calls; write issues_preview.md")
    p.add_argument("--limit", type=int, help="create at most this many new issues")
    p.add_argument("--done", default="", help="comma-separated IDs to close as completed")
    p.add_argument("--only", default="", help="comma-separated IDs or gates (e.g. G0,G1) to import")
    p.add_argument("--exclude-category", default="", help="comma-separated categories to skip")
    p.add_argument("--no-assign", action="store_true", help="leave issues unassigned")
    p.add_argument("--handle", action="append", default=[], metavar="NAME=LOGIN",
                   help="override a GitHub handle, e.g. --handle Ricky=ricky-p")
    p.add_argument("--delay", type=float, default=2.0, help="seconds between writes (default 2)")
    p.add_argument("--out", type=Path, default=HERE / "issue_numbers.csv")
    p.add_argument("--api-url", default=API_URL, help=argparse.SUPPRESS)
    args = p.parse_args(argv)

    board = read_board(args.board)
    tasks = [normalise(r) for r in board["tasks"]]
    tasks = [t for t in tasks if t["id"]]
    ids = [t["id"] for t in tasks]
    dupes = sorted({i for i in ids if ids.count(i) > 1})
    if dupes:
        sys.exit(f"duplicate Task IDs on the board: {', '.join(dupes)}")
    unknown = sorted({(t["id"], d) for t in tasks for d in t["depends"] if d not in ids})
    if unknown:
        print("warning: unknown dependency IDs (left as text): "
              + ", ".join(f"{t}->{d}" for t, d in unknown))
    for t in tasks:
        t["labels"] = t["labels"] or derive_labels(t)

    handles = {**HANDLES, **board["handles"]}
    for item in args.handle:
        name, _, login = item.partition("=")
        handles[name.strip()] = login.strip()

    excluded = {c.strip().lower() for c in args.exclude_category.split(",") if c.strip()}
    only = {c.strip() for c in args.only.split(",") if c.strip()}
    selected = [t for t in tasks if t["category"].lower() not in excluded
                and (not only or t["id"] in only or t["gate"] in only)]
    selected = topological(selected)
    done_ids = {d.strip() for d in args.done.split(",") if d.strip()}
    imported_on = dt.date.today().strftime("%d %b %Y").lstrip("0")

    print(f"board: {len(tasks)} tasks; selected {len(selected)}"
          + (f" (excluded categories: {', '.join(sorted(excluded))})" if excluded else ""))

    if args.dry_run:
        numbers = {t["id"]: "?" for t in selected}
        preview = HERE / "issues_preview.md"
        with preview.open("w", encoding="utf-8") as fh:
            for t in selected:
                fh.write(f"# [{t['id']}] {t['title']}\n\n")
                fh.write(f"labels: {', '.join(t['labels'])} | milestone: {t['gate']} | "
                         f"assignee: {'-' if args.no_assign else handles.get(t['owner'], '-')}\n\n")
                fh.write(issue_body(t, handles, numbers, imported_on) + "\n---\n\n")
        gates: dict[str, int] = {}
        owners: dict[str, int] = {}
        for t in selected:
            gates[t["gate"]] = gates.get(t["gate"], 0) + 1
            owners[t["owner"]] = owners.get(t["owner"], 0) + 1
        print("per gate : " + ", ".join(f"{g} {n}" for g, n in sorted(gates.items())))
        print("per owner: " + ", ".join(f"{o} {n}" for o, n in sorted(owners.items())))
        print(f"labels   : {len({lab for t in selected for lab in t['labels']})}")
        missing = sorted({t["owner"] for t in selected if t["owner"] not in handles})
        if missing and not args.no_assign:
            print(f"no handle (left unassigned): {', '.join(missing)}")
        print(f"dry run: nothing sent. All issue bodies written to {preview}")
        return 0

    gh = GitHub(args.repo, find_token(), args.api_url, args.delay)

    # Existing issues (open and closed; pull requests excluded) -> skip those IDs.
    numbers: dict[str, int] = {}
    for item in gh.paged("/issues?state=all"):
        if "pull_request" in item:
            continue
        m = TITLE_ID.match(item.get("title", ""))
        if m:
            numbers.setdefault(m.group(1), item["number"])
    for t in tasks:  # ids already recorded on the board
        if t["issue"].lstrip("#").isdigit():
            numbers.setdefault(t["id"], int(t["issue"].lstrip("#")))
    todo = [t for t in selected if t["id"] not in numbers]
    existing = len(selected) - len(todo)
    if args.limit is not None:
        todo = todo[: args.limit]
    print(f"{args.repo}: {existing} already exist, creating {len(todo)}")
    if not todo and not done_ids:
        return 0

    # Labels
    have = {lab["name"].lower() for lab in gh.paged("/labels")}
    needed = sorted({lab for t in todo for lab in t["labels"]})
    for name in needed:
        if name.lower() in have:
            continue
        colour, desc = board["labels"].get(name) or (
            next((c for k, c in LABEL_COLOURS.items() if name.startswith(k)), "0F766E"),
            name.replace(":", ": ").replace("-", " "))
        gh.call("POST", "/labels", {"name": name, "color": colour.lstrip("#"),
                                    "description": desc[:100]})
        have.add(name.lower())
        print(f"  label {name}")

    # Milestones, one per gate, matched by the gate code at the start of the title
    milestones = {}
    for ms in gh.paged("/milestones?state=all"):
        code = ms["title"].split()[0].rstrip(":") if ms["title"].strip() else ""
        milestones.setdefault(code, ms["number"])
    for gate in sorted({t["gate"] for t in todo if t["gate"]}):
        if gate in milestones:
            continue
        title, due, exit_criteria = board["milestones"].get(gate) or (
            GATES.get(gate, ("", None))[0], _date(GATES.get(gate, ("", None))[1]), "")
        payload = {"title": f"{gate} {title}".strip(), "description": exit_criteria}
        if due:
            payload["due_on"] = f"{due.isoformat()}T12:00:00Z"  # noon UTC: same date in Sydney and the US
        ms, _ = gh.call("POST", "/milestones", payload)
        milestones[gate] = ms["number"]
        print(f"  milestone {payload['title']}")

    # Which handles can be assigned (repo collaborators). GET /assignees/{login}
    # answers 204 when assignable and 404 otherwise.
    assignable: dict[str, bool] = {}
    if not args.no_assign:
        for login in sorted({handles[t["owner"]] for t in todo if t["owner"] in handles}):
            found, _ = gh.call("GET", f"/assignees/{urllib.parse.quote(login)}", ok_missing=True)
            assignable[login] = found is not MISSING
            if not assignable[login]:
                print(f"  warning: {login} cannot be assigned in {args.repo}; "
                      "their issues stay unassigned (owner is still named in the body)")

    new_file = not args.out.exists()
    with args.out.open("a", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        if new_file:
            writer.writerow(["Task ID", "GitHub Issue #", "URL"])
        for n, t in enumerate(todo, 1):
            payload = {"title": f"[{t['id']}] {t['title']}",
                       "body": issue_body(t, handles, numbers, imported_on),
                       "labels": t["labels"]}
            if t["gate"] in milestones:
                payload["milestone"] = milestones[t["gate"]]
            login = handles.get(t["owner"])
            if login and assignable.get(login):
                payload["assignees"] = [login]
            issue, _ = gh.call("POST", "/issues", payload)
            numbers[t["id"]] = issue["number"]
            writer.writerow([t["id"], issue["number"], issue.get("html_url", "")])
            fh.flush()
            print(f"  [{n}/{len(todo)}] #{issue['number']} {t['id']}")
            status = t["status"].lower()
            if status == "done":
                done_ids.add(t["id"])
            elif status == "cut":
                gh.call("PATCH", f"/issues/{issue['number']}",
                        {"state": "closed", "state_reason": "not_planned"})

    for tid in sorted(done_ids):
        if tid not in numbers:
            print(f"  warning: --done {tid}: no issue found")
            continue
        gh.call("PATCH", f"/issues/{numbers[tid]}", {"state": "closed", "state_reason": "completed"})
        print(f"  closed #{numbers[tid]} {tid} as completed")

    print(f"done: {len(todo)} created, {gh.writes} writes. Issue numbers appended to {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
